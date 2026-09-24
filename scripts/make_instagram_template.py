"""
Generates docs/Brand_Gadari_Kishor_Kumar_Instagram_TEMPLATE.xlsx: a blank
workbook with the exact section layout scripts/import_politicians.py expects,
pre-filled with headers, labels and the 30 daily date columns for
Aug 20 - Sep 18, 2026, so it drops straight into the existing import
pipeline once the data cells are filled in.

Usage:
    python3 scripts/make_instagram_template.py
"""

import os

import openpyxl
from openpyxl.styles import Font, PatternFill

OUT_PATH = os.path.join(
    os.path.dirname(__file__), "..", "docs", "Brand_Gadari_Kishor_Kumar_Instagram_TEMPLATE.xlsx"
)

DATES = [
    "20 August 2026", "21 August 2026", "22 August 2026", "23 August 2026", "24 August 2026",
    "25 August 2026", "26 August 2026", "27 August 2026", "28 August 2026", "29 August 2026",
    "30 August 2026", "31 August 2026", "01 September 2026", "02 September 2026", "03 September 2026",
    "04 September 2026", "05 September 2026", "06 September 2026", "07 September 2026", "08 September 2026",
    "09 September 2026", "10 September 2026", "11 September 2026", "12 September 2026", "13 September 2026",
    "14 September 2026", "15 September 2026", "16 September 2026", "17 September 2026", "18 September 2026",
]

SECTION_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
SECTION_FONT = Font(bold=True)
HEADER_FONT = Font(bold=True, color="666666")
NOTE_FONT = Font(italic=True, color="94A3B8")


def section_title(ws, row, title):
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL


def sub_header(ws, row, values):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=1 + i, value=v)
        c.font = HEADER_FONT


def daily_header_row(ws, row):
    ws.cell(row=row, column=1, value="")
    for i, d in enumerate(DATES):
        ws.cell(row=row, column=2 + i, value=d)


def blank_daily_row(ws, row, label):
    ws.cell(row=row, column=1, value=label)
    # data cells (col 2..31) intentionally left blank for daily counts


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gadari Kishor Kumar Instagram"

    ws.cell(row=1, column=1, value="Socialinsider (manual template)")
    ws.cell(row=2, column=1, value="Gadari Kishor Kumar Brand - Instagram")
    ws.cell(row=2, column=10, value="August,20 2026 - September,18 2026")

    ws.cell(row=4, column=1, value=(
        "Fill in every blank cell below with real, verifiable Instagram data for "
        "@dr.gadarikishorekumar covering Aug 20 - Sep 18, 2026. Leave a cell blank only if "
        "that metric genuinely isn't available (e.g. Instagram doesn't expose it publicly) - "
        "do not guess or estimate. This file follows the exact layout "
        "scripts/import_politicians.py expects, so once filled in it can be merged into the "
        "dashboard the same way as the other three politicians' reports."
    )).font = NOTE_FONT

    # --- Brand profiles ---
    section_title(ws, 6, "Brand profiles")
    sub_header(ws, 8, ["", "Platform", "Followers", "Engagement", "Eng. rate per followers", "Posts"])
    ws.cell(row=9, column=1, value="dr.gadarikishorekumar")
    ws.cell(row=9, column=2, value="ig")

    # --- KPIs ---
    section_title(ws, 12, "Key Performance Indicators")
    sub_header(ws, 14, ["KPI", "Current", "Previous", "Change %"])
    for i, kpi in enumerate(["Posts", "Engagement", "Fans Count", "Avg Posts / Day", "Avg Engagement", "Total Video Views", "Total Impressions"]):
        ws.cell(row=15 + i, column=1, value=kpi)

    # --- Audience: followers by platform ---
    section_title(ws, 24, "Audience - Followers by platform")
    sub_header(ws, 26, ["", "Followers", "Followers diff. percent vs previous period"])
    ws.cell(row=27, column=1, value="Instagram")

    # --- Audience: followers net diff ---
    section_title(ws, 30, "Audience - Followers by platform net difference")
    sub_header(ws, 32, ["", "Followers diff", "Followers growth percent vs previous period"])
    ws.cell(row=33, column=1, value="Instagram")

    # --- Posts distribution (daily) ---
    section_title(ws, 36, "Posts Distribution Across Channels")
    daily_header_row(ws, 38)
    blank_daily_row(ws, 39, "Instagram")

    # --- Total posts by platform ---
    section_title(ws, 42, "Total Posts by Platform")
    sub_header(ws, 44, ["", "Posts", "Vs previous period"])
    ws.cell(row=45, column=1, value="Instagram")

    # --- Engagement distribution (daily) ---
    section_title(ws, 48, "Engagement Distribution Across Channels")
    daily_header_row(ws, 50)
    blank_daily_row(ws, 51, "Instagram")

    # --- Total engagement by platform ---
    section_title(ws, 54, "Total Engagement by Platform")
    sub_header(ws, 56, ["", "Engagement", "Vs previous period"])
    ws.cell(row=57, column=1, value="Instagram")

    # --- Avg engagement rate by platform ---
    section_title(ws, 60, "Avg. Engagement Rate by Platform")
    sub_header(ws, 62, ["", "Engagement rate", "Vs previous period"])
    ws.cell(row=63, column=1, value="Instagram")

    # --- Views distribution (daily) ---
    section_title(ws, 66, "Views Distribution Across Channels")
    daily_header_row(ws, 68)
    blank_daily_row(ws, 69, "Instagram")

    # --- Total views by platform ---
    section_title(ws, 72, "Total Views Across Channels")
    sub_header(ws, 74, ["", "Views", "Views diff. percent vs previous period"])
    ws.cell(row=75, column=1, value="Instagram")

    # --- Video views distribution (daily) ---
    section_title(ws, 78, "Video Views Distribution Across Channels")
    daily_header_row(ws, 80)
    blank_daily_row(ws, 81, "Instagram")

    # --- Video views totals by platform ---
    section_title(ws, 84, "Video Views Totals Across Channels")
    sub_header(ws, 86, ["", "Video Views", "Video views diff. percent vs previous period"])
    ws.cell(row=87, column=1, value="Instagram")

    # --- Top posts (up to 5) ---
    section_title(ws, 90, "Top posts")
    sub_header(ws, 92, ["Page", "Link", "Date", "Type", "", "", "", "", "Engagement", "Engagement rate"])
    for i in range(5):
        ws.cell(row=93 + i, column=1, value="dr.gadarikishorekumar")

    # --- PDF-supplement fields (hand-transcribed, matching data/pdf_supplement.json shape) ---
    pdf_row = 100
    section_title(ws, pdf_row, "PDF supplement fields (not read by the importer - transcribe by hand into data/pdf_supplement.json)")
    ws.cell(row=pdf_row + 2, column=1, value="Brand Comments (current, change %)")
    ws.cell(row=pdf_row + 3, column=1, value="Brand Likes (current, change %)")
    ws.cell(row=pdf_row + 4, column=1, value="Brand Growth of Followers (current, change %)")
    ws.cell(row=pdf_row + 6, column=1, value="Content pillars (name, posts, engagement, avg eng. rate, top channel) - up to 4 rows")
    ws.cell(row=pdf_row + 12, column=1, value="Earned Media Value - Instagram (value, avg value)")

    # Column widths for readability
    ws.column_dimensions["A"].width = 32
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 16

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote template to {OUT_PATH}")


if __name__ == "__main__":
    main()
