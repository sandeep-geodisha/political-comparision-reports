"""
Parses the Socialinsider brand-report .xlsx exports in docs/ into a single
JSON file (data/politicians.json) consumed by the Next.js dashboard.

Usage:
    python3 scripts/import_politicians.py

Re-run this whenever a new or updated .xlsx report is dropped into docs/.
Each source file must follow the same fixed-row-layout export template
used by the three existing reports (see docs/*.xlsx).
"""

import glob
import json
import os
import re

import openpyxl

DOCS_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")
OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "politicians.json")

PLATFORM_NAMES = {"fb": "facebook", "ig": "instagram", "tw": "twitter"}


def slugify(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return s


def cell(ws, row, col):
    return ws.cell(row=row, column=col).value


def row_values(ws, row, max_col):
    return [ws.cell(row=row, column=c).value for c in range(1, max_col + 1)]


def parse_pct(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 4)
    s = str(v).replace("%", "").strip()
    try:
        return round(float(s), 4)
    except ValueError:
        return None


def parse_platform_table(ws, header_row, max_col=3):
    """Rows shaped: ['', colA_label, colB_label], then one row per platform."""
    labels = row_values(ws, header_row, max_col + 1)[1:]
    out = {}
    r = header_row + 1
    while True:
        name = cell(ws, r, 1)
        if not name or not isinstance(name, str):
            break
        vals = row_values(ws, r, max_col + 1)[1:]

        def coerce(v):
            if isinstance(v, str):
                try:
                    return round(float(v.replace("%", "").strip()), 4)
                except ValueError:
                    return v
            return v

        out[name.strip().lower()] = {labels[i]: coerce(vals[i]) for i in range(len(labels))}
        r += 1
    return out


def parse_daily_series(ws, header_row):
    """Rows shaped: ['', date1, date2, ...], then one row per platform of daily values."""
    dates = []
    c = 2
    while True:
        v = cell(ws, header_row, c)
        if v is None:
            break
        dates.append(str(v))
        c += 1
    series = {}
    r = header_row + 1
    while True:
        name = cell(ws, r, 1)
        if not name or not isinstance(name, str):
            break
        vals = [cell(ws, r, c) or 0 for c in range(2, 2 + len(dates))]
        series[name.strip().lower()] = vals
        r += 1
    return dates, series


def find_section_row(ws, title, max_row):
    for r in range(1, max_row + 1):
        if str(cell(ws, r, 1) or "").strip() == title:
            return r
    raise ValueError(f"Section '{title}' not found")


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    max_row = ws.max_row

    date_range = cell(ws, 2, 10) or cell(ws, 2, 9) or ""

    # --- Brand profiles ---
    profiles_header = find_section_row(ws, "Brand profiles", max_row) + 2
    profiles = []
    r = profiles_header + 1
    while True:
        handle = cell(ws, r, 1)
        platform_code = cell(ws, r, 2)
        if not platform_code:
            break
        profiles.append(
            {
                "handle": handle,
                "platform": PLATFORM_NAMES.get(str(platform_code).strip().lower(), platform_code),
                "followers": cell(ws, r, 3),
                "engagement": cell(ws, r, 4),
                "engagementRatePerFollower": cell(ws, r, 5),
                "posts": cell(ws, r, 6),
            }
        )
        r += 1

    # --- KPIs ---
    kpi_header = find_section_row(ws, "Key Performance Indicators", max_row) + 2
    kpis = {}
    r = kpi_header + 1
    while True:
        name = cell(ws, r, 1)
        if not name:
            break
        kpis[name.strip()] = {
            "current": cell(ws, r, 2),
            "previous": cell(ws, r, 3),
            "changePct": parse_pct(cell(ws, r, 4)),
        }
        r += 1

    # --- Audience: followers by platform ---
    aud_header = find_section_row(ws, "Audience - Followers by platform", max_row) + 2
    followers_by_platform = parse_platform_table(ws, aud_header, max_col=2)

    # --- Audience: followers net diff ---
    aud_diff_header = find_section_row(ws, "Audience - Followers by platform net difference", max_row) + 2
    followers_diff_by_platform = parse_platform_table(ws, aud_diff_header, max_col=2)

    # --- Posts distribution (daily) ---
    posts_dist_header = find_section_row(ws, "Posts Distribution Across Channels", max_row) + 2
    posts_dates, posts_daily = parse_daily_series(ws, posts_dist_header)

    # --- Total posts by platform ---
    total_posts_header = find_section_row(ws, "Total Posts by Platform", max_row) + 2
    total_posts_by_platform = parse_platform_table(ws, total_posts_header, max_col=2)

    # --- Engagement distribution (daily) ---
    eng_dist_header = find_section_row(ws, "Engagement Distribution Across Channels", max_row) + 2
    _, engagement_daily = parse_daily_series(ws, eng_dist_header)

    # --- Total engagement by platform ---
    total_eng_header = find_section_row(ws, "Total Engagement by Platform", max_row) + 2
    total_engagement_by_platform = parse_platform_table(ws, total_eng_header, max_col=2)

    # --- Avg engagement rate by platform ---
    eng_rate_header = find_section_row(ws, "Avg. Engagement Rate by Platform", max_row) + 2
    engagement_rate_by_platform = parse_platform_table(ws, eng_rate_header, max_col=2)

    # --- Views distribution (daily) ---
    views_dist_header = find_section_row(ws, "Views Distribution Across Channels", max_row) + 2
    _, views_daily = parse_daily_series(ws, views_dist_header)

    # --- Total views by platform ---
    total_views_header = find_section_row(ws, "Total Views Across Channels", max_row) + 2
    total_views_by_platform = parse_platform_table(ws, total_views_header, max_col=2)

    # --- Video views distribution (daily) ---
    vviews_dist_header = find_section_row(ws, "Video Views Distribution Across Channels", max_row) + 2
    _, video_views_daily = parse_daily_series(ws, vviews_dist_header)

    # --- Video views totals by platform ---
    total_vviews_header = find_section_row(ws, "Video Views Totals Across Channels", max_row) + 2
    total_video_views_by_platform = parse_platform_table(ws, total_vviews_header, max_col=2)

    # --- Top posts ---
    top_posts_header = find_section_row(ws, "Top posts", max_row) + 2
    top_posts = []
    r = top_posts_header + 1
    while True:
        page = cell(ws, r, 1)
        link = cell(ws, r, 2)
        if not page and not link:
            break
        top_posts.append(
            {
                "page": page,
                "link": link,
                "date": cell(ws, r, 3),
                "type": cell(ws, r, 4),
                "engagement": cell(ws, r, 9),
                "engagementRate": parse_pct(cell(ws, r, 10)),
            }
        )
        r += 1

    return {
        "dateRange": str(date_range),
        "profiles": profiles,
        "kpis": kpis,
        "audience": {
            "followersByPlatform": followers_by_platform,
            "followersDiffByPlatform": followers_diff_by_platform,
        },
        "posts": {
            "dates": posts_dates,
            "dailyByPlatform": posts_daily,
            "totalByPlatform": total_posts_by_platform,
        },
        "engagement": {
            "dailyByPlatform": engagement_daily,
            "totalByPlatform": total_engagement_by_platform,
            "rateByPlatform": engagement_rate_by_platform,
        },
        "views": {
            "dailyByPlatform": views_daily,
            "totalByPlatform": total_views_by_platform,
        },
        "videoViews": {
            "dailyByPlatform": video_views_daily,
            "totalByPlatform": total_video_views_by_platform,
        },
        "topPosts": top_posts,
    }


# Filenames look like: Brand_<Politician_Name>_<start>_<end>_<hash> 1.xlsx
FILENAME_RE = re.compile(r"^Brand_(.+?)_\d{1,2}_[A-Za-z]{3}_\d{4}_\d{1,2}_[A-Za-z]{3}_\d{4}_")


def name_from_filename(filename):
    base = os.path.basename(filename)
    m = FILENAME_RE.match(base)
    raw = m.group(1) if m else base
    return raw.replace("_", " ").strip()


def main():
    files = sorted(glob.glob(os.path.join(DOCS_DIR, "*.xlsx")))
    if not files:
        raise SystemExit(f"No .xlsx files found in {DOCS_DIR}")

    politicians = []
    for path in files:
        name = name_from_filename(path)
        parsed = parse_workbook(path)
        politicians.append(
            {
                "id": slugify(name),
                "name": name,
                "sourceFile": os.path.basename(path),
                **parsed,
            }
        )
        print(f"Parsed {name!r} from {os.path.basename(path)}")

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w") as f:
        json.dump({"politicians": politicians}, f, indent=2)

    print(f"\nWrote {len(politicians)} politicians to {OUT_PATH}")


if __name__ == "__main__":
    main()
